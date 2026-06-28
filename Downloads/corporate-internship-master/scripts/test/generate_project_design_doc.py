from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "企业实习WMS与AI库存预测系统设计文档.docx"

FONT_CN = "Microsoft YaHei"
FONT_BODY = "Calibri"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(90, 98, 112)
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9DEE7"
WHITE = "FFFFFF"


def set_run_font(run, name=FONT_BODY, east_asia=FONT_CN, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def paragraph_border_bottom(paragraph, color="2E74B5", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_text(doc, text, style=None, bold=False, italic=False, color=None, size=None, align=None, before=None, after=None):
    paragraph = doc.add_paragraph(style=style)
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(item, tuple):
            label, detail = item
            r1 = p.add_run(label)
            set_run_font(r1, bold=True)
            r2 = p.add_run(detail)
            set_run_font(r2)
        else:
            r = p.add_run(item)
            set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    set_table_borders(table, color="CDD6E3", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    set_table_borders(table)
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], header_fill)
        p = header_cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return table


def add_cover(doc):
    add_text(doc, "项目设计文档", bold=True, size=23, color=NAVY, after=4)
    add_text(doc, "企业实习 WMS 与 AI 库存预测系统", size=15, color=GRAY, after=14)

    metadata = [
        ("项目路径", r"D:\code\corporate-internship-master\corporate-internship"),
        ("文档类型", "系统设计说明书"),
        ("版本", "V1.0"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("适用范围", "前端管理后台、Spring Boot 后端、FastAPI AI 预测服务、MySQL 数据库"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)
    rule = doc.add_paragraph()
    paragraph_border_bottom(rule, color="2E74B5")
    add_callout(
        doc,
        "文档说明",
        "本文档基于当前项目源码、数据库表结构说明、AI 服务 README 和项目启动记录整理，重点描述系统架构、模块边界、数据模型、接口设计、核心业务流程和部署测试方案。"
    )


def add_toc(doc):
    doc.add_page_break()
    add_text(doc, "目录", style="Heading 1")
    toc_items = [
        "1. 项目概述",
        "2. 总体架构设计",
        "3. 技术栈与运行环境",
        "4. 功能模块设计",
        "5. 数据库设计",
        "6. 后端接口设计",
        "7. AI 需求预测与智能预警设计",
        "8. 前端交互与移动访问设计",
        "9. 核心业务流程",
        "10. 部署启动设计",
        "11. 测试与验证设计",
        "12. 风险与后续优化",
        "附录 A. 常用启动与测试指令",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_body(doc):
    doc.add_page_break()
    add_text(doc, "1. 项目概述", style="Heading 1")
    add_text(
        doc,
        "本项目是一个面向仓储场景的后台管理系统，核心目标是实现入库、出库、看板、库存追溯、转包、封存和 AI 需求预测的一体化管理。系统由 Vue 前端、Spring Boot 后端、FastAPI AI 服务和 MySQL 数据库组成，前端通过代理访问普通业务接口和 AI 预测接口。"
    )
    add_bullets(doc, [
        ("业务范围：", "供应商管理、零件管理、入库单、出库单、扫码入库、扫码出库、实时库存、库存追溯、转包、封存、AI 需求预测。"),
        ("设计目标：", "保证库存数量口径统一，支持看板二维码流转，保留完整库存流水，并通过 AI 服务给出缺货、呆滞和波动预警。"),
        ("当前重点：", "AI 需求预测统一为“出库量预测”，趋势图保留历史入库量、出库量、当前库存和未来出库需求预测线。"),
    ])

    add_text(doc, "2. 总体架构设计", style="Heading 1")
    add_text(doc, "系统采用前后端分离与 AI 微服务并行的结构。Spring Boot 负责 WMS 业务数据写入和查询，FastAPI 直接读取同一 MySQL 数据库做预测分析，Vue 前端通过 Vite 代理统一调用。")
    add_table(
        doc,
        ["层级", "组成", "职责", "端口/访问方式"],
        [
            ["表现层", "Vue 3 + Vite + Element Plus", "后台页面、路由、表格、表单、扫码、打印、ECharts 趋势图", "https://localhost:5173"],
            ["业务服务层", "Spring Boot backend-demo", "登录、菜单、WMS 入库出库、看板、库存、追溯、转包、封存等接口", "http://localhost:8080/api"],
            ["AI 服务层", "FastAPI ai-predict-service", "读取库存流水与库存汇总，生成预测趋势和预警看板", "http://localhost:8001"],
            ["数据层", "MySQL backend_db", "保存业务主数据、单据、看板、库存、追溯和用户数据", "localhost:3306"],
        ],
        [1.0, 1.7, 2.8, 1.0],
        header_fill=BLUE_GRAY,
    )
    add_callout(
        doc,
        "接口访问边界",
        "前端普通业务接口使用 `/api` 代理到 Spring Boot；AI 预测接口使用 `/ai-api` 代理到 FastAPI。这样手机或其他局域网设备访问前端时，不需要把 AI 服务地址写死为某台电脑的 IP。"
    )

    add_text(doc, "3. 技术栈与运行环境", style="Heading 1")
    add_table(
        doc,
        ["子系统", "主要技术", "关键配置"],
        [
            ["前端", "Vue 3.5、Vite 8、Element Plus 2.14、Vue Router、Axios、ECharts、QRCode、ZXing", "开发服务 `5173`，启用 HTTPS，`host=0.0.0.0` 支持局域网访问"],
            ["后端", "Java 17、Spring Boot 4.0.6、Spring WebMVC、Spring JDBC、MySQL Connector/J", "`server.port=8080`，上下文路径 `/api`，数据库 `backend_db`"],
            ["AI 服务", "FastAPI、Uvicorn、SQLAlchemy、PyMySQL、pandas、numpy、scikit-learn、matplotlib", "通过 `MYSQL_HOST`、`MYSQL_PORT`、`MYSQL_USER`、`MYSQL_PASSWORD`、`MYSQL_DATABASE` 连接数据库"],
            ["数据库", "MySQL 8.4，本地开发库 `backend_db`", "字符集建议 `utf8mb4`，本地开发账号当前使用 `root/root`"],
        ],
        [1.0, 3.0, 2.5],
    )

    add_text(doc, "4. 功能模块设计", style="Heading 1")
    add_table(
        doc,
        ["模块", "页面/入口", "设计说明"],
        [
            ["基础资料", "供应商管理、零件管理", "维护供应商和零件基础信息，为入库单、出库单和库存统计提供主数据。"],
            ["入库管理", "入库单列表、创建入库单、扫码入库", "创建入库单和看板，扫码确认后写入当前库存与入库追溯流水。"],
            ["出库管理", "出库单列表、创建出库单、扫码出库", "按出库单校验看板，确认出库后扣减或移除当前库存，并写入出库追溯流水。"],
            ["库存总览", "实时库存总览", "按零件、仓库、供应商汇总当前有效库存，过滤已封存看板。"],
            ["库存追溯", "库存追溯、看板生命周期", "基于 `inventory_trace` 查询入库、出库、转包、封存、解封等完整流水。"],
            ["转包管理", "转包列表、拆分、合并", "支持看板拆分和合并，并保留父子看板关系和追溯记录。"],
            ["封存管理", "封存看板管理", "封存后库存不参与普通可用库存统计，支持解封和历史查询。"],
            ["AI 需求预测", "AI需求预测", "显示三类预警汇总、零件监控列表和 7 天出入库趋势预测图。"],
        ],
        [1.15, 1.55, 3.8],
    )

    add_text(doc, "5. 数据库设计", style="Heading 1")
    add_text(doc, "数据库围绕看板和库存流水展开。`kanban` 是库存流转载体，`current_inventory` 表示当前在库状态，`inventory_trace` 保存所有业务动作，是 AI 预测和追溯分析的主要数据来源。")
    add_table(
        doc,
        ["表名", "类型", "作用"],
        [
            ["kanban", "核心业务表", "保存看板号、关联单据、零件、仓库、数量、状态、封存和转包关系。"],
            ["current_inventory", "核心业务表", "保存当前在库看板库存，入库新增，出库删除或扣减，封存时状态变化。"],
            ["inventory_trace", "核心业务表", "保存 INBOUND、OUTBOUND、TRANSFER_SPLIT、TRANSFER_MERGE、SEAL、UNSEAL 等流水。"],
            ["inbound_order / inbound_order_detail", "单据表", "保存入库单主表和明细行，记录预期数量、已收数量和箱数。"],
            ["outbound_order / outbound_order_detail", "单据表", "保存出库单主表和明细行，记录预期出库和已发数量。"],
            ["transfer_order", "业务表", "保存转包拆分与合并记录。"],
            ["seal_history", "业务表", "保存看板封存与解封历史。"],
            ["part / supplier / warehouse / user", "基础表", "保存零件、供应商、仓库和登录用户。"],
        ],
        [1.75, 1.15, 3.6],
    )
    add_bullets(doc, [
        ("库存口径：", "`current_inventory.quantity > 0` 且看板未封存时，才计入实时可用库存。"),
        ("追溯口径：", "AI 趋势分析按最近 7 天 `inventory_trace` 聚合入库量和出库量。"),
        ("看板状态：", "`pending` 表示待入库，`stored` 表示已入库，`outbound` 表示已出库。"),
        ("单据状态：", "入库和出库单均按 `pending -> scanning -> completed` 流转。"),
    ])

    add_text(doc, "6. 后端接口设计", style="Heading 1")
    add_text(doc, "Spring Boot 后端使用 `/api` 作为上下文路径，WMS 业务接口集中在 `WmsController` 的 `/wms` 路径下；登录和菜单分别由 `LoginController`、`MenuController` 提供。")
    add_table(
        doc,
        ["接口组", "典型路径", "说明"],
        [
            ["登录认证", "POST `/api/login`、POST `/api/logout`", "处理后台登录、退出和会话状态。"],
            ["菜单", "GET `/api/menus`", "返回左侧菜单和路由所需数据。"],
            ["基础资料", "GET/POST/DELETE `/api/wms/supplier/*`、`/api/wms/part/*`", "供应商和零件的查询、保存和删除。"],
            ["入库", "`/api/wms/inbound-order/*`、POST `/api/wms/scan/inbound`", "入库单创建、编辑、打印看板、扫码入库。"],
            ["出库", "`/api/wms/outbound-order/*`、POST `/api/wms/scan/outbound`", "出库单创建、编辑、打印看板、扫码出库。"],
            ["库存与追溯", "GET `/api/wms/inventory/current`、`/api/wms/trace/list`、`/api/wms/trace/by-kanban/{kanbanNo}`", "实时库存查询、库存流水和看板生命周期查询。"],
            ["转包与封存", "POST `/api/wms/transfer/split`、POST `/api/wms/transfer/merge`、POST `/api/wms/kanban/seal`", "看板拆分、合并、封存和解封。"],
        ],
        [1.4, 2.6, 2.5],
    )

    add_text(doc, "7. AI 需求预测与智能预警设计", style="Heading 1")
    add_text(doc, "AI 服务作为独立 FastAPI 微服务运行，直接读取 MySQL 中的库存和追溯数据。当前设计中，“AI 需求预测”明确指未来出库量预测，不再预测入库量或总流量。")
    add_table(
        doc,
        ["组件", "文件", "职责"],
        [
            ["数据服务", "services/data_service.py", "读取零件库存汇总、零件列表和最近 7 天库存流水。"],
            ["特征工程", "models/feature_engineer.py", "构造日期特征、滞后特征、移动平均、趋势和波动统计。"],
            ["预测器", "models/predictor.py", "集成移动平均、线性回归、随机森林，并对上升、下降、起伏序列做模式识别预测。"],
            ["异常检测", "models/anomaly_detector.py", "识别波动异常，要求多次方向反转和足够振幅，避免单日突增误判。"],
            ["预警服务", "services/warning_service.py", "生成缺货、呆滞和波动预警，并给出处理建议。"],
            ["趋势服务", "services/trend_service.py", "生成趋势图 JSON 数据，返回历史入库、历史出库、库存和预测出库量。"],
        ],
        [1.1, 1.9, 3.5],
    )
    add_bullets(doc, [
        ("缺货预警：", "预测 7 天出库需求大于当前库存时触发；库存低于安全库存时给出中风险提醒。"),
        ("呆滞预警：", "库存可维持时间过长或近期无出库时触发，用于识别周转慢的库存。"),
        ("波动预警：", "近 7 天出库或流量出现多次方向反转并有明显振幅时触发。"),
        ("趋势图预测模式：", "`predictionMeta.mode` 可返回 `ensemble`、`rising`、`falling`、`fluctuating` 或 `none`。"),
    ])
    add_table(
        doc,
        ["测试零件", "用途", "期望表现"],
        [
            ["AITEST-001", "缺货预警", "低库存、近 7 天持续出库，触发缺货卡片。"],
            ["AITEST-002", "呆滞预警", "高库存、近 7 天无出库，触发呆滞卡片。"],
            ["AITEST-003", "波动预警 + 起伏预测", "出库量低高交替，红色预测线呈有起有伏。"],
            ["AITEST-004", "上升预测", "近 7 天出库量持续上升，红色预测线继续上升。"],
            ["AITEST-005", "下降预测", "近 7 天出库量持续下降，红色预测线继续下降。"],
        ],
        [1.2, 1.8, 3.5],
        header_fill=BLUE_GRAY,
    )

    add_text(doc, "8. 前端交互与移动访问设计", style="Heading 1")
    add_text(doc, "前端使用 Vue Router 按后台菜单组织页面。`AdminLayout` 负责左侧菜单、顶部标签页和内容区域，WMS 页面集中在 `frontend-admin/src/views/wms`。")
    add_bullets(doc, [
        ("接口封装：", "`src/api/http.js` 创建 Axios 客户端，普通业务接口默认使用 `/api`。"),
        ("AI 接口：", "`src/api/aiPredict.js` 通过 `/ai-api` 访问 FastAPI 服务，由 Vite 代理转发到 `localhost:8001`。"),
        ("移动访问：", "Vite 使用 `host=0.0.0.0` 和 HTTPS，手机需与电脑在同一局域网，通过电脑局域网 IP 和 5173 端口访问。"),
        ("可视化：", "AI 趋势弹窗使用 ECharts 绘制入库量、出库量、当前库存和 AI 需求预测线。"),
    ])

    add_text(doc, "9. 核心业务流程", style="Heading 1")
    add_text(doc, "9.1 入库流程", style="Heading 2")
    add_numbered(doc, [
        "用户创建入库单，填写供应商、仓库、零件、预期数量和包装容量。",
        "系统生成或打印入库看板二维码，每个看板对应零件、数量、箱号和单据。",
        "扫码入库时校验看板合法性和剩余可入库数量。",
        "确认入库后，写入 `current_inventory`，更新入库单已收数量，并写入 `inventory_trace` 的 INBOUND 流水。",
        "全部看板入库完成后，入库单状态变为 completed。",
    ])
    add_text(doc, "9.2 出库流程", style="Heading 2")
    add_numbered(doc, [
        "用户创建出库单，填写客户、仓库、零件和预期出库数量。",
        "系统根据出库需求生成出库看板，或通过已有库存看板扫码出库。",
        "扫码出库时校验看板是否存在、是否已封存、是否属于当前出库零件和仓库。",
        "确认出库后，从 `current_inventory` 删除或扣减库存，更新出库单已发数量，并写入 OUTBOUND 追溯流水。",
        "出库数量达到预期后，出库单状态变为 completed。",
    ])
    add_text(doc, "9.3 AI 预测流程", style="Heading 2")
    add_numbered(doc, [
        "前端进入 AI 需求预测页面，请求 `/ai-api/api/dashboard/summary` 和 `/ai-api/api/parts/all`。",
        "AI 服务读取当前库存汇总和最近 7 天追溯流水。",
        "按零件聚合历史入库量、出库量和当前库存，生成预警列表。",
        "用户点击趋势时，请求 `/ai-api/api/predict/trend/{part_code}`。",
        "AI 服务返回历史 7 天数据和未来 7 天出库需求预测，前端用 ECharts 渲染趋势图。",
    ])

    add_text(doc, "10. 部署启动设计", style="Heading 1")
    add_table(
        doc,
        ["启动对象", "目录", "推荐命令"],
        [
            ["MySQL", "项目根目录", 'start "mysql-local" /B "C:\\Program Files\\MySQL\\MySQL Server 8.4\\bin\\mysqld.exe" --basedir="C:\\Program Files\\MySQL\\MySQL Server 8.4" --datadir="%CD%\\.local-mysql\\data" --port=3306 --bind-address=127.0.0.1 --log-error="%CD%\\.local-mysql\\logs\\mysql.err"'],
            ["后端", "backend-demo", "java -jar target\\backend-demo-0.0.1-SNAPSHOT.jar"],
            ["AI 服务", "ai-predict-service", ".venv\\Scripts\\python.exe -m uvicorn app:app --host 0.0.0.0 --port 8001"],
            ["前端", "frontend-admin", "npm run dev"],
        ],
        [1.0, 1.25, 4.25],
    )
    add_callout(
        doc,
        "启动顺序建议",
        "先启动 MySQL，再启动 Spring Boot 后端，然后启动 FastAPI AI 服务，最后启动 Vite 前端。前端页面登录后即可访问 WMS 与 AI 预测功能。"
    )

    add_text(doc, "11. 测试与验证设计", style="Heading 1")
    add_table(
        doc,
        ["测试类型", "命令/位置", "验证目标"],
        [
            ["前端单元测试", "`frontend-admin`: `npm test`", "路由、代理、AI 页面、扫码出库、打印等源码级约束。"],
            ["AI 合约测试", "`ai-predict-service`: `.venv\\Scripts\\python.exe -m unittest tests.test_ai_inventory_contracts`", "库存汇总、预警、异常检测和预测形态。"],
            ["脚本测试", "项目根目录：`python -m unittest scripts.tests.test_ai_trend_seed_scripts`", "测试数据 seed/clear 范围和用例完整性。"],
            ["接口验证", "`/api/dashboard/summary`、`/api/predict/trend/AITEST-*`", "三类预警和上升、下降、起伏预测结果。"],
            ["数据库验证", "执行 `clear-ai-full-test-data.sql` 和 `seed-ai-full-test-data.sql`", "确保测试数据可重复生成、可清理。"],
        ],
        [1.25, 2.55, 2.7],
    )

    add_text(doc, "12. 风险与后续优化", style="Heading 1")
    add_bullets(doc, [
        ("数据量风险：", "当前 AI 预测主要基于最近 7 天数据，测试环境数据较少时预测更偏规则和模式识别。后续可增加更长时间窗口和真实业务样本。"),
        ("认证安全：", "当前后台以本地开发为主，后续可增强 Token、角色权限、密码加密和接口鉴权。"),
        ("服务部署：", "当前三服务独立启动，后续可引入统一启动脚本、Windows Service 或容器编排。"),
        ("预测效果：", "当前模型适合演示和趋势参考，生产场景建议加入节假日、订单计划、安全库存策略和模型评估报表。"),
        ("数据一致性：", "入库、出库、转包、封存需要持续保证 `kanban`、`current_inventory` 和 `inventory_trace` 三者同步。"),
    ])

    doc.add_page_break()
    add_text(doc, "附录 A. 常用启动与测试指令", style="Heading 1")
    add_text(doc, "A.1 CMD 启动顺序", style="Heading 2")
    add_table(
        doc,
        ["步骤", "命令"],
        [
            ["1", r"cd /d D:\code\corporate-internship-master\corporate-internship"],
            ["2", 'start "mysql-local" /B "C:\\Program Files\\MySQL\\MySQL Server 8.4\\bin\\mysqld.exe" --basedir="C:\\Program Files\\MySQL\\MySQL Server 8.4" --datadir="%CD%\\.local-mysql\\data" --port=3306 --bind-address=127.0.0.1 --log-error="%CD%\\.local-mysql\\logs\\mysql.err"'],
            ["3", r"cd backend-demo && java -jar target\backend-demo-0.0.1-SNAPSHOT.jar"],
            ["4", r"cd ..\ai-predict-service && set MYSQL_HOST=localhost && set MYSQL_PORT=3306 && set MYSQL_USER=root && set MYSQL_PASSWORD=root && set MYSQL_DATABASE=backend_db && .venv\Scripts\python.exe -m uvicorn app:app --host 0.0.0.0 --port 8001"],
            ["5", r"cd ..\frontend-admin && npm run dev"],
        ],
        [0.55, 5.95],
    )
    add_text(doc, "A.2 AI 测试数据脚本", style="Heading 2")
    add_table(
        doc,
        ["用途", "命令"],
        [
            ["清理 AI 测试数据", '"C:\\Program Files\\MySQL\\MySQL Server 8.4\\bin\\mysql.exe" --default-character-set=utf8mb4 -uroot -proot backend_db < scripts\\clear-ai-full-test-data.sql'],
            ["重新生成 AI 测试数据", '"C:\\Program Files\\MySQL\\MySQL Server 8.4\\bin\\mysql.exe" --default-character-set=utf8mb4 -uroot -proot backend_db < scripts\\seed-ai-full-test-data.sql'],
        ],
        [1.4, 5.1],
    )


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("企业实习 WMS 与 AI 库存预测系统 | 设计文档")
    set_run_font(r, size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = footer.add_run("Confidential - Local Project Design")
    set_run_font(r2, size=9, color=GRAY)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    set_header_footer(doc)
    add_cover(doc)
    add_toc(doc)
    add_body(doc)
    doc.core_properties.title = "企业实习 WMS 与 AI 库存预测系统设计文档"
    doc.core_properties.subject = "WMS, AI Demand Prediction, System Design"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "Generated from local project source and documentation."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
